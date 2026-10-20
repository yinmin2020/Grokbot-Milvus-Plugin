"""Ingest documents and retrieve cited chunks through MilvusClient.

Run with uv run --project /path/to/plugin /path/to/plugin/scripts/milvus_kb.py.
Document IDs use resolved paths by default, or an explicit namespace + source root.
Retries verify the complete chunk set before declaring a document unchanged.
"""

import argparse
import hashlib
import json
import os
import re
import sys
import time
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from mvbot_common import (
    Embedder,
    add_embed_args,
    add_milvus_args,
    check_write,
    collection_meta,
    connect,
    default_prefixes,
    ensure_loaded,
    env,
    literal,
    print_json,
    query_rows,
    redact_uri,
    resolve_embedding,
    validate_schema,
    validate_vectors,
)
from mvbot_search import search

SCHEMA_TAG = "mvbot-kb/v1"
CHUNKER_VERSION = "1"
DEFAULT_CHUNK, DEFAULT_OVERLAP = (
    int(os.environ.get("KB_CHUNK_SIZE", "500")),
    int(os.environ.get("KB_CHUNK_OVERLAP", "80")),
)
EXTS = {".md", ".markdown", ".txt", ".pdf", ".docx"}
ANALYZER = {"tokenizer": "jieba", "filter": ["lowercase", "cnalphanumonly"]}  # zh + en
OUT_FIELDS = ["pk", "doc_id", "source", "title", "heading", "page", "chunk_no", "lang", "text"]
CJK = re.compile(r"[\u3400-\u9fff\uf900-\ufaff]")


# ------------------------------------------------------------------ parsing
def _paras(text):
    return [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]


def parse_markdown(text, fallback_title):
    blocks, stack, buf, in_code, title = [], [], [], False, None

    def flush():
        if buf:
            body = "\n".join(buf).strip()
            if body:
                blocks.append((" > ".join(h for _, h in stack), None, body))
            buf.clear()

    for line in text.splitlines():
        if line.strip().startswith("```"):
            in_code = not in_code
            buf.append(line)
            continue
        m = None if in_code else re.match(r"^(#{1,6})\s+(.+?)\s*#*\s*$", line)
        if m:
            flush()
            level, h = len(m.group(1)), m.group(2).strip()
            if level == 1 and title is None:
                title = h
            stack[:] = [(level_value, x) for level_value, x in stack if level_value < level] + [(level, h)]
        elif not line.strip() and not in_code:
            flush()
        else:
            buf.append(line)
    flush()
    return title or fallback_title, blocks


def parse_txt(text, fallback_title):
    ps = _paras(text)
    if len(ps) > 1 and "\n" not in ps[0] and len(ps[0]) <= 120:  # short first line = title
        return ps[0], [("", None, p) for p in ps[1:]]
    return fallback_title, [("", None, p) for p in ps]


def parse_pdf(path, fallback_title):
    from pypdf import PdfReader

    r = PdfReader(path)
    title = (r.metadata.title if r.metadata and r.metadata.title else None) or fallback_title
    blocks = []
    for i, page in enumerate(r.pages, 1):
        t = page.extract_text() or ""
        ps = _paras(t) or [level_value.strip() for level_value in t.splitlines() if level_value.strip()]
        blocks += [("", i, re.sub(r"[ \t]*\n[ \t]*", "\n", p)) for p in ps]
    return title, blocks


def parse_docx(path, fallback_title):
    import docx

    d = docx.Document(path)
    title = (d.core_properties.title or "").strip() or None
    blocks, stack = [], []
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name if p.style is not None else "") or ""
        m = re.match(r"^(Heading|标题)\s*(\d)", style)
        if style == "Title":
            title = title or t
            continue
        if m:
            level = int(m.group(2))
            if level == 1 and title is None:
                title = t
            stack[:] = [(level_value, x) for level_value, x in stack if level_value < level] + [(level, t)]
            continue
        blocks.append((" > ".join(h for _, h in stack), None, t))
    for ti, table in enumerate(d.tables, 1):
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        if rows:
            blocks.append((f"table {ti}", None, "\n".join(rows)))
    return title or fallback_title, blocks


def parse_file(path):
    ext = os.path.splitext(path)[1].lower()
    fallback = os.path.splitext(os.path.basename(path))[0]
    if ext == ".pdf":
        return parse_pdf(path, fallback)
    if ext == ".docx":
        return parse_docx(path, fallback)
    text = open(path, encoding="utf-8", errors="replace").read()
    if ext in (".md", ".markdown"):
        return parse_markdown(text, fallback)
    return parse_txt(text, fallback)


# ------------------------------------------------------------------ chunking
_SENT = re.compile(r"(?<=[。！？!?；;])|(?<=[.!?])\s+")


def _split_long(text, size):
    sents = [s for s in _SENT.split(text) if s and s.strip()]
    out, cur = [], ""
    for s in sents:
        while len(s) > size:  # hard split very long sentences
            if cur:
                out.append(cur)
                cur = ""
            out.append(s[:size])
            s = s[size:]
        if len(cur) + len(s) > size and cur:
            out.append(cur)
            cur = ""
        sep = (
            "" if (not cur or cur.endswith((" ", "\n")) or CJK.search(cur[-1:]) or CJK.search(s[:1])) else " "
        )
        cur += sep + s
    if cur.strip():
        out.append(cur)
    return out


def _tail(text, overlap):
    if overlap <= 0 or len(text) <= overlap:
        return text if overlap > 0 else ""
    t = text[-overlap:]
    m = re.search(r"[。！？!?；;.]\s*", t)  # start the overlap at a sentence boundary when possible
    return t[m.end() :] if m and m.end() < len(t) - 10 else t


def chunk_blocks(blocks, size, overlap):
    """blocks: [(heading, page, text)] -> [(heading, page, chunk_text)]; never crosses a section."""
    chunks, i = [], 0
    while i < len(blocks):
        heading, page = blocks[i][0], blocks[i][1]
        sec = []
        while i < len(blocks) and blocks[i][0] == heading and blocks[i][1] == page:
            sec.append(blocks[i][2])
            i += 1
        pieces = []
        for p in sec:
            lim = max(50, size - overlap)  # leave room for the overlap prefix
            pieces += _split_long(p, lim) if len(p) > lim else [p]
        cur = ""
        for p in pieces:
            if cur and len(cur) + len(p) + 1 > size:
                chunks.append((heading, page, cur))
                tail = _tail(cur, overlap)
                cur = (tail + "\n" if tail else "") + p
            else:
                cur = (cur + "\n" + p) if cur else p
        if cur.strip():
            chunks.append((heading, page, cur))
    return chunks


def lang_of(text):
    if not text:
        return "und"
    return "zh" if len(CJK.findall(text)) / max(1, len(text)) > 0.15 else "en"


def clip_bytes(s, n):
    b = s.encode("utf-8")
    return s if len(b) <= n else b[:n].decode("utf-8", errors="ignore")


# ------------------------------------------------------------------ collection
def read_meta(c, col):
    from pymilvus import DataType

    schema, meta = collection_meta(c, col, SCHEMA_TAG)
    fields = {
        name: DataType.VARCHAR
        for name in ["pk", "doc_id", "doc_hash", "source", "title", "heading", "lang", "text"]
    }
    fields.update({name: DataType.INT64 for name in ["page", "chunk_no", "updated_at"]})
    fields.update(dense=DataType.FLOAT_VECTOR, sparse=DataType.SPARSE_FLOAT_VECTOR)
    validate_schema(schema, fields, meta["dim"])
    if not any(
        f.get("input_field_names") == ["text"] and f.get("output_field_names") == ["sparse"]
        for f in schema.get("functions", [])
    ):
        raise ValueError("Expected a text-to-sparse BM25 function")
    return schema, meta


def create_collection(c, col, emb, dim, prefixes, a):
    from pymilvus import DataType, Function, FunctionType

    size = a.chunk_size if a.chunk_size is not None else DEFAULT_CHUNK
    overlap = a.chunk_overlap if a.chunk_overlap is not None else DEFAULT_OVERLAP
    if size < 100 or not 0 <= overlap < size:
        raise ValueError("Require chunk_size >= 100 and 0 <= chunk_overlap < chunk_size")
    meta = {
        "schema": SCHEMA_TAG,
        "embed_provider": emb.provider,
        "embed_model": emb.model,
        "dim": dim,
        "metric": "COSINE",
        "query_prefix": prefixes["query"],
        "doc_prefix": prefixes["doc"],
        "min_score": prefixes.get("min_score", 0.5),
        "chunk_size": size,
        "chunk_overlap": overlap,
        "analyzer": ANALYZER,
        "created_at": int(time.time()),
    }
    s = c.create_schema(
        auto_id=False, enable_dynamic_field=False, description=json.dumps(meta, ensure_ascii=False)
    )
    s.add_field("pk", DataType.VARCHAR, is_primary=True, max_length=160)
    s.add_field("doc_id", DataType.VARCHAR, max_length=40)
    s.add_field("doc_hash", DataType.VARCHAR, max_length=64)
    s.add_field("source", DataType.VARCHAR, max_length=1024)
    s.add_field("title", DataType.VARCHAR, max_length=512)
    s.add_field("heading", DataType.VARCHAR, max_length=1024)
    s.add_field("page", DataType.INT64)
    s.add_field("chunk_no", DataType.INT64)
    s.add_field("lang", DataType.VARCHAR, max_length=8)
    s.add_field("updated_at", DataType.INT64)
    s.add_field(
        "text",
        DataType.VARCHAR,
        max_length=65535,
        enable_analyzer=True,
        enable_match=True,
        analyzer_params=ANALYZER,
    )
    s.add_field("dense", DataType.FLOAT_VECTOR, dim=dim)
    s.add_field("sparse", DataType.SPARSE_FLOAT_VECTOR)
    s.add_function(
        Function(
            name="text_bm25",
            function_type=FunctionType.BM25,
            input_field_names=["text"],
            output_field_names=["sparse"],
        )
    )
    ip = c.prepare_index_params()
    ip.add_index("dense", index_type="AUTOINDEX", metric_type="COSINE")
    ip.add_index(
        "sparse",
        index_type="SPARSE_INVERTED_INDEX",
        metric_type="BM25",
        params={"inverted_index_algo": "DAAT_MAXSCORE"},
    )
    ip.add_index("doc_id", index_type="INVERTED")
    ip.add_index("source", index_type="INVERTED")
    check_write(col)
    c.create_collection(
        col,
        schema=s,
        index_params=ip,
        consistency_level="Strong",
        properties={"mvbot.config": json.dumps(meta, ensure_ascii=False)},
    )
    return meta


def resolve_embedder(c, a, create=False):
    """Return (embedder, meta). Uses the model recorded in the collection unless overridden;
    refuses to mix models. If create=True and the collection is missing, creates it."""
    col = a.collection
    exists = c.has_collection(col)
    meta = None
    if exists:
        _, meta = read_meta(c, col)
        for key in ("query_prefix", "doc_prefix"):
            override = getattr(a, key, None)
            if override is not None and override != meta[key]:
                raise ValueError(f"{key} differs from collection configuration; use a new collection")
        emb = resolve_embedding(a, meta)
        return emb, meta
    if not create:
        sys.exit(f"ERROR: collection '{col}' does not exist. Run `init` or `ingest` first.")
    emb = Embedder(a.embed_provider, a.embed_endpoint, a.embed_model, a.embed_batch, a.embed_timeout)
    dim = emb.dim()
    prefixes = default_prefixes(emb.model)
    if a.query_prefix is not None:
        prefixes["query"] = a.query_prefix
    if a.doc_prefix is not None:
        prefixes["doc"] = a.doc_prefix
    meta = create_collection(c, col, emb, dim, prefixes, a)
    print(f"[created] collection '{col}' dense dim={dim} model={emb.provider}:{emb.model} + BM25 sparse")
    return emb, meta


# ------------------------------------------------------------------ commands
def iter_files(paths):
    for p in paths:
        if os.path.isdir(p):
            for root, dirs, files in os.walk(p):
                dirs[:] = sorted(d for d in dirs if not d.startswith("."))
                for f in sorted(files):
                    if os.path.splitext(f)[1].lower() in EXTS and not f.startswith((".", "~$")):
                        yield os.path.join(root, f)
        elif os.path.isfile(p):
            if os.path.splitext(p)[1].lower() in EXTS:
                yield p
            else:
                print(f"[skip] unsupported file type: {p}")
        else:
            print(f"[skip] not found: {p}")


def source_identity(path, a):
    resolved = Path(path).resolve()
    if bool(a.namespace) != bool(a.source_root):
        raise ValueError("Use --namespace and --source-root together for portable document IDs")
    if a.source_root:
        try:
            name = resolved.relative_to(Path(a.source_root).resolve()).as_posix()
        except ValueError as exc:
            raise ValueError(f"File is outside --source-root: {path}") from exc
        identity = [a.namespace, name]
    else:
        name = resolved.as_posix()
        identity = ["absolute-path", name]
    source = (a.source_prefix.rstrip("/") + "/" + name) if a.source_prefix else name
    if len(source.encode()) > 1024:
        raise ValueError("Source exceeds 1024 UTF-8 bytes; use --namespace and --source-root")
    return source, hashlib.sha256(json.dumps(identity).encode()).hexdigest()[:40]


def cmd_init(c, a):
    check_write(a.collection)
    if c.has_collection(a.collection):
        _, meta = resolve_embedder(c, a)
        print_json({"collection": a.collection, "status": "exists", "config": meta})
        return
    resolve_embedder(c, a, create=True)


def cmd_ingest(c, a):
    check_write(a.collection)
    exists = c.has_collection(a.collection)
    if a.dry_run:
        meta = (
            read_meta(c, a.collection)[1]
            if exists
            else {"chunk_size": DEFAULT_CHUNK, "chunk_overlap": DEFAULT_OVERLAP, "doc_prefix": ""}
        )
        emb = None
    else:
        emb, meta = resolve_embedder(c, a, create=True)
        ensure_loaded(c, a.collection)
    size = a.chunk_size if a.chunk_size is not None else meta["chunk_size"]
    overlap = a.chunk_overlap if a.chunk_overlap is not None else meta["chunk_overlap"]
    if size < 100 or not 0 <= overlap < size:
        raise ValueError("Require chunk_size >= 100 and 0 <= chunk_overlap < chunk_size")
    files = list(iter_files(a.paths))
    if not files:
        raise ValueError("No supported files found")
    summary = {"added": 0, "updated": 0, "unchanged": 0, "chunks_written": 0, "stale_deleted": 0}
    for path in files:
        src, doc_id = source_identity(path, a)
        raw = Path(path).read_bytes()
        doc_hash = hashlib.sha256(raw + f"|{size}|{overlap}|{CHUNKER_VERSION}|{src}".encode()).hexdigest()
        title, blocks = parse_file(path)
        chunks = chunk_blocks(blocks, size, overlap)
        if not chunks:
            raise ValueError(
                f"No extractable text in {src}; existing chunks retained. "
                "Use delete-doc to remove them explicitly, or OCR a scanned PDF first"
            )
        flt = f"doc_id == {literal(doc_id)}"
        old = list(query_rows(c, a.collection, flt, ["pk", "doc_hash"])) if emb else []
        expected = {f"{doc_id}#{n:05d}" for n in range(len(chunks))}
        complete = {r["pk"] for r in old} == expected and all(r["doc_hash"] == doc_hash for r in old)
        if complete and not a.force:
            summary["unchanged"] += 1
            print(f"[unchanged] {src}")
            continue
        now = int(time.time())
        rows, inputs = [], []
        for n, (heading, page, body) in enumerate(chunks):
            text = clip_bytes((f"[{heading}]\n" if heading else "") + body, 65000)
            rows.append(
                {
                    "pk": f"{doc_id}#{n:05d}",
                    "doc_id": doc_id,
                    "doc_hash": doc_hash,
                    "source": src,
                    "title": clip_bytes(title, 512),
                    "heading": clip_bytes(heading, 1024),
                    "page": page or 0,
                    "chunk_no": n,
                    "lang": lang_of(body),
                    "updated_at": now,
                    "text": text,
                }
            )
            inputs.append(meta["doc_prefix"].replace("{title}", title or "none") + text)
        if a.dry_run:
            print(f"[dry-run] {src}: title={title!r} chunks={len(rows)} (not written)")
            continue
        vecs = emb.embed(inputs)
        validate_vectors(vecs, meta["dim"], len(rows))
        for row, vector in zip(rows, vecs):
            row["dense"] = vector
        for i in range(0, len(rows), 64):
            c.upsert(a.collection, rows[i : i + 64])
        stale = [r for r in old if r["pk"] not in expected]
        if stale:
            c.delete(a.collection, filter=flt + f" and chunk_no >= {len(rows)}")
        kind = "updated" if old else "added"
        summary[kind] += 1
        summary["chunks_written"] += len(rows)
        summary["stale_deleted"] += len(stale)
        print(f"[{kind}] {src}: {len(rows)} chunks, {len(stale)} stale chunks deleted")
    print("summary:", json.dumps(summary))


def _hit(h):
    """pymilvus 2.x/3.x compatible (pk, distance, entity) for a search hit."""
    e = h.get("entity") if isinstance(h, dict) and isinstance(h.get("entity"), dict) else None
    e = e or (h.entity if hasattr(h, "entity") and isinstance(getattr(h, "entity"), dict) else {}) or h
    pk = getattr(h, "id", None) or getattr(h, "pk", None)
    if pk is None:
        for k in ("id", "pk"):
            try:
                pk = h[k]
                break
            except (KeyError, TypeError):
                pass
    if pk is None:
        pk = e.get("pk")
    dist = getattr(h, "distance", None)
    if dist is None:
        dist = h["distance"]
    return pk, float(dist), e


def cmd_query(c, a):
    _, meta = read_meta(c, a.collection)
    if a.min_score is not None and a.mode != "dense":
        raise ValueError("--min-score is a cosine threshold for --mode dense only")
    qvec = None
    if a.mode != "bm25":
        emb = resolve_embedding(a, meta)
        qvec = emb.embed([meta["query_prefix"] + a.question])[0]
        validate_vectors([qvec], meta["dim"])
    hits = search(
        c,
        a.collection,
        mode=a.mode,
        text=a.question,
        vector=qvec,
        filter=a.filter,
        top_k=a.top_k,
        candidates=max(a.candidates, a.top_k),
        ranker=a.ranker,
        rrf_k=a.rrf_k,
        weights=tuple(float(x) for x in a.weights.split(",")),
        output_fields=OUT_FIELDS,
    )
    threshold = a.min_score if a.min_score is not None else float(meta.get("min_score", 0.5))
    results = []
    for h in hits:
        pk, score, e = _hit(h)
        if a.mode == "dense" and score < threshold:
            continue
        results.append(
            {
                "cite": len(results) + 1,
                "pk": pk,
                "score": score,
                "dense_cosine": score if a.mode == "dense" else None,
                **{key: e.get(key) for key in OUT_FIELDS if key != "pk"},
            }
        )
    out = {
        "question": a.question,
        "collection": a.collection,
        "mode": a.mode,
        "found": bool(results),
        "min_score": threshold if a.mode == "dense" else None,
        "score_kind": {"dense": "cosine", "bm25": "bm25", "hybrid": a.ranker}[a.mode],
        "results": results,
    }
    # found means retrieved candidates; relevance and answerability still require reading the text.
    if a.json:
        print_json(out)
        return
    print(f"Q: {a.question}\nmode={a.mode} found={out['found']} score={out['score_kind']}")
    for r in results:
        print(f"\n[{r['cite']}] {r['source']} > {r['heading'] or ''} (page {r['page'] or '-'})")
        print(f"    pk={r['pk']} score={r['score']:.4f}")
        print(r["text"] if a.full else r["text"][:300])


def _docs(c, col, flt=""):
    docs = {}
    for r in query_rows(c, col, flt, ["doc_id", "source", "title", "doc_hash", "updated_at"]):
        d = docs.setdefault(
            r["doc_id"],
            {
                "doc_id": r["doc_id"],
                "source": r["source"],
                "title": r["title"],
                "doc_hash": r["doc_hash"][:12],
                "updated_at": r["updated_at"],
                "chunks": 0,
            },
        )
        d["chunks"] += 1
    return sorted(docs.values(), key=lambda d: d["source"])


def cmd_info(c, a):
    d, meta = read_meta(c, a.collection)
    ensure_loaded(c, a.collection)
    rows = c.query(a.collection, filter="", output_fields=["count(*)"], consistency_level="Strong")[0][
        "count(*)"
    ]
    docs = _docs(c, a.collection)
    print_json(
        {
            "collection": a.collection,
            "uri": redact_uri(a.uri),
            "kb_meta": meta,
            "row_count": rows,
            "documents": len(docs),
            "load_state": str(c.get_load_state(a.collection).get("state")),
            "fields": [f["name"] for f in d["fields"]],
            "functions": [f.get("name") for f in d.get("functions", [])],
        }
    )


def cmd_list_docs(c, a):
    read_meta(c, a.collection)
    ensure_loaded(c, a.collection)
    docs = _docs(c, a.collection)
    for d in docs:
        print(f"{d['doc_id']}  chunks={d['chunks']:<4} hash={d['doc_hash']}  {d['source']}  ({d['title']})")
    print(f"total documents: {len(docs)}")


def cmd_delete_doc(c, a):
    check_write(a.collection)
    read_meta(c, a.collection)
    if bool(a.source) == bool(a.doc_id):
        raise ValueError("Pass exactly one of --source or --doc-id")
    flt = f"doc_id == {literal(a.doc_id)}" if a.doc_id else f"source == {literal(a.source)}"
    ensure_loaded(c, a.collection)
    n = c.query(a.collection, filter=flt, output_fields=["count(*)"], consistency_level="Strong")[0][
        "count(*)"
    ]
    if not a.execute:
        print(
            f"[dry-run] would delete {n} chunks from '{a.collection}' where {flt}. "
            f"Re-run with --execute --confirm {a.collection}"
        )
        return
    if a.confirm != a.collection:
        sys.exit(f"REFUSED: --confirm must equal the collection name '{a.collection}'")
    c.delete(a.collection, filter=flt)
    print(f"[deleted] {n} chunks where {flt}")


def build_parser():
    default_col = env("KB_COLLECTION", "kb_docs")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    def common(p, embed=True):
        add_milvus_args(p)
        p.add_argument("--collection", default=default_col)
        if embed:
            add_embed_args(p)
        return p

    for name in ("init", "ingest"):
        p = common(sub.add_parser(name))
        p.add_argument(
            "--chunk-size",
            type=int,
            default=None,
            help="characters per chunk (default: collection setting, new collections 500 or env KB_CHUNK_SIZE)",
        )
        p.add_argument(
            "--chunk-overlap",
            type=int,
            default=None,
            help="default: collection setting, new collections 80 or env KB_CHUNK_OVERLAP",
        )
        p.add_argument("--query-prefix", default=None, help="override model task prefix for queries")
        p.add_argument(
            "--doc-prefix", default=None, help="override model task prefix for documents ({title} allowed)"
        )
        if name == "ingest":
            p.add_argument("paths", nargs="+")
            p.add_argument("--source-prefix", default="", help="e.g. a URL base stored in `source`")
            p.add_argument("--namespace", default="", help="stable project identity, requires --source-root")
            p.add_argument(
                "--source-root", default="", help="root for portable source paths, requires --namespace"
            )
            p.add_argument("--force", action="store_true", help="re-embed even if unchanged")
            p.add_argument("--dry-run", action="store_true", help="parse + chunk only, write nothing")
    p = common(sub.add_parser("query"))
    p.add_argument("question")
    p.add_argument("--mode", choices=["hybrid", "dense", "bm25"], default="hybrid")
    p.add_argument("--top-k", type=int, default=5)
    p.add_argument("--candidates", type=int, default=20)
    p.add_argument("--rrf-k", type=int, default=60)
    p.add_argument("--filter", default="", help='Milvus filter, e.g. lang == "zh" or source like "docs/%%"')
    p.add_argument(
        "--min-score",
        type=float,
        default=None,
        help="dense mode only: discard hits below this cosine threshold",
    )
    p.add_argument("--ranker", choices=["rrf", "weighted"], default="rrf")
    p.add_argument("--weights", default="0.7,0.3", help="dense,bm25 weights for --ranker weighted")
    p.add_argument("--json", action="store_true")
    p.add_argument("--full", action="store_true", help="print full chunk text")
    common(sub.add_parser("info"), embed=False)
    common(sub.add_parser("list-docs"), embed=False)
    p = common(sub.add_parser("delete-doc"), embed=False)
    p.add_argument("--source")
    p.add_argument("--doc-id")
    p.add_argument("--execute", action="store_true")
    p.add_argument("--confirm", default="")
    return ap


def main():
    a = build_parser().parse_args()
    c = connect(a)
    try:
        if a.cmd not in ("init", "ingest") and not c.has_collection(a.collection):
            raise ValueError(f"Collection '{a.collection}' does not exist")
        globals()["cmd_" + a.cmd.replace("-", "_")](c, a)
    finally:
        c.close()


if __name__ == "__main__":
    try:
        main()
    except ValueError as exc:
        sys.exit(f"ERROR: {exc}")
