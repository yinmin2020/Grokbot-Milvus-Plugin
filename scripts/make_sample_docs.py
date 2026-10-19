# /// script
# requires-python = ">=3.10"
# dependencies = ["python-docx>=1.1", "fpdf2>=2.7"]
# ///
"""Generate the binary sample documents (docx + pdf) used by the Knowledge Desk end-to-end test.

  uv run scripts/make_sample_docs.py [out_dir]     # default: samples/kb
"""
import os, sys

out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "samples", "kb")
os.makedirs(out, exist_ok=True)

import docx
d = docx.Document()
d.core_properties.title = "Milvus 运维手册"
d.add_heading("Milvus 运维手册", level=1)
d.add_heading("健康检查", level=2)
d.add_paragraph("Milvus 在 HTTP 端口（默认 9091）上提供 /healthz 健康检查接口和 /metrics 监控指标接口，监控指标采用 Prometheus 格式，可以接入 Grafana 面板。")
d.add_heading("Compaction 压缩", level=2)
d.add_paragraph("删除数据后，Milvus 不会立即回收空间，而是先写入删除记录（delete log）。Compaction 会合并小 Segment 并清理已删除的数据。可以调用 compact 接口手动触发，再用 get_compaction_state 查询进度。")
d.add_heading("备份与恢复", level=2)
d.add_paragraph("官方推荐使用 milvus-backup 工具进行备份。备份前先执行 flush，把增长中的 Segment 封存到对象存储。恢复时建议使用新的 Collection 名称后缀，验证行数一致后再切换别名。")
t = d.add_table(rows=3, cols=2)
for r, (k, v) in enumerate([("端口", "用途"), ("19530", "gRPC / SDK 访问"), ("9091", "HTTP：healthz 与 metrics")]):
    t.cell(r, 0).text, t.cell(r, 1).text = k, v
d.save(os.path.join(out, "ops_manual_zh.docx"))

from fpdf import FPDF
pdf = FPDF(); pdf.set_title("Milvus Partitions and Memory")
pages = [
    ("Partitions", "A partition is a subset of a collection. Every collection has a _default partition. "
     "You can create up to 1024 partitions per collection by default. Searching with partition_names restricts the "
     "search to those partitions and reduces the amount of data scanned. For multi-tenant data with many tenants, "
     "prefer a partition key field instead of manual partitions."),
    ("Load and release", "A collection must be loaded into memory before it can be searched or queried. "
     "release_collection frees the memory on query nodes; the data stays in object storage and can be loaded again. "
     "An index can only be dropped after the collection has been released."),
]
for title, body in pages:
    pdf.add_page(); pdf.set_font("Helvetica", "B", 16); pdf.cell(0, 10, title, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=11); pdf.multi_cell(0, 6, body)
pdf.output(os.path.join(out, "partitions_memory_en.pdf"))
print("wrote", os.listdir(out))
